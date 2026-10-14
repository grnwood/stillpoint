from pathlib import Path
import os
import subprocess
import time
import zipfile

from docx import Document
import pytest
from pptx import Presentation
from pptx.util import Inches

from sp.app.folder_navigator import documents


def _sample_docx(path: Path) -> None:
    document = Document()
    document.add_heading("Quarterly Review", level=1)
    paragraph = document.add_paragraph()
    paragraph.add_run("Important").bold = True
    paragraph.add_run(" details")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Amount"
    table.cell(1, 0).text = "Revenue"
    table.cell(1, 1).text = "$42"
    document.add_paragraph("After the table")
    document.save(path)


def test_docx_falls_back_to_portable_structured_html(tmp_path, monkeypatch):
    path = tmp_path / "review.docx"
    _sample_docx(path)
    monkeypatch.setattr(documents, "_libreoffice_executable", lambda: None)

    preview = documents.build_docx_preview(path)

    assert preview.backend == "python-docx"
    assert preview.pdf_path is None
    assert "Simplified preview" in preview.notice
    assert "<h1>Quarterly Review</h1>" in preview.html
    assert "<strong>Important</strong> details" in preview.html
    assert "<table>" in preview.html
    assert preview.html.index("Quarterly Review") < preview.html.index("Revenue")
    assert preview.html.index("Revenue") < preview.html.index("After the table")


@pytest.mark.parametrize(
    ("suffix", "builder_name", "limit_name"),
    [
        (".docx", "build_docx_preview", "MAX_DOCX_UNCOMPRESSED_BYTES"),
        (".pptx", "build_pptx_preview", "MAX_PPTX_UNCOMPRESSED_BYTES"),
    ],
)
def test_office_preview_rejects_archive_expansion_bombs(
        tmp_path, monkeypatch, suffix, builder_name, limit_name):
    path = tmp_path / f"bomb{suffix}"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", b"x" * 256)
    monkeypatch.setattr(documents, limit_name, 128)

    with pytest.raises(ValueError, match="expands beyond"):
        getattr(documents, builder_name)(path)


def test_docx_libreoffice_conversion_is_fingerprint_cached(tmp_path, monkeypatch):
    path = tmp_path / "review.docx"
    _sample_docx(path)
    cache = tmp_path / "cache"
    monkeypatch.setattr(documents, "_libreoffice_executable", lambda: "/usr/bin/soffice")
    calls = []

    def fake_run(command, **kwargs):
        calls.append((command, kwargs))
        output = Path(command[command.index("--outdir") + 1])
        (output / "review.pdf").write_bytes(b"%PDF-1.4\npreview\n")
        return subprocess.CompletedProcess(command, 0, b"", b"")

    monkeypatch.setattr(documents.subprocess, "run", fake_run)

    first = documents.build_docx_preview(path, cache)
    second = documents.build_docx_preview(path, cache)

    assert first.backend == "libreoffice"
    assert first.pdf_path == second.pdf_path
    assert first.pdf_path.read_bytes().startswith(b"%PDF")
    assert "Quarterly Review" in first.html
    assert len(calls) == 1
    assert calls[0][1]["timeout"] == 45


def test_pptx_falls_back_to_portable_slide_preview(tmp_path, monkeypatch):
    path = tmp_path / "briefing.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Project Briefing"
    text_box = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(5), Inches(1))
    paragraph = text_box.text_frame.paragraphs[0]
    paragraph.add_run().text = "Delivery status"
    paragraph.runs[0].font.bold = True
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(5), Inches(1.5)).table
    table.cell(0, 0).text = "Workstream"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "Preview"
    table.cell(1, 1).text = "Ready"
    slide.notes_slide.notes_text_frame.text = "Discuss the preview fallback."
    presentation.save(path)
    monkeypatch.setattr(documents, "_libreoffice_executable", lambda: None)

    preview = documents.build_pptx_preview(path)

    assert preview.backend == "python-pptx"
    assert preview.pdf_path is None
    assert "slide-layout fidelity" in preview.notice
    assert "Project Briefing" in preview.html
    assert "<strong>Delivery status</strong>" in preview.html
    assert "Workstream" in preview.html
    assert "Speaker notes" in preview.html
    assert "Discuss the preview fallback." in preview.html


@pytest.mark.parametrize(
    ("suffix", "builder_name", "backend"),
    [
        (".docx", "build_docx_preview", "python-docx"),
        (".pptx", "build_pptx_preview", "python-pptx"),
    ],
)
def test_folder_navigator_displays_portable_office_preview(
        tmp_path, monkeypatch, qapp, suffix, builder_name, backend):
    from PySide6.QtCore import Qt
    from PySide6.QtWidgets import QPushButton, QTextBrowser
    from sp.app.folder_navigator.window import Window

    path = tmp_path / f"portable{suffix}"
    path.write_bytes(b"placeholder")
    monkeypatch.setattr(Path, "home", lambda: tmp_path)
    monkeypatch.setattr(
        documents,
        builder_name,
        lambda _path: documents.DocxPreview(
            backend=backend,
            html="<h1>Portable preview</h1>",
            notice="Simplified preview notice",
        ),
    )
    window = Window(tmp_path)
    window.open_file(path, pinned=True)

    deadline = time.monotonic() + 2
    while not window.active_tab().findChildren(QTextBrowser):
        assert time.monotonic() < deadline
        qapp.processEvents()
        time.sleep(.01)

    viewer = window.active_tab().viewer
    browser = viewer.findChild(QTextBrowser)
    assert browser.toPlainText() == "Portable preview"
    assert browser.textInteractionFlags() & Qt.TextSelectableByMouse
    assert any(
        button.text() == "Open in Default Application"
        for button in viewer.findChildren(QPushButton)
    )
    menu = window._create_preview_context_menu(window.active_tab(), browser)
    assert "Reveal in Folder" in [action.text() for action in menu.actions()]
    window.close()


def test_office_preview_cache_is_bounded_by_file_count(tmp_path, monkeypatch):
    cache = tmp_path / "cache"
    cache.mkdir()
    files = [cache / f"{index}.pdf" for index in range(3)]
    for index, path in enumerate(files):
        path.write_bytes(b"pdf")
        os.utime(path, (100 + index, 100 + index))
    monkeypatch.setattr(documents, "MAX_OFFICE_CACHE_FILES", 2)
    monkeypatch.setattr(documents, "MAX_OFFICE_CACHE_BYTES", 1024)
    monkeypatch.setattr(documents, "MAX_OFFICE_CACHE_AGE_SECONDS", 10**12)

    documents._trim_office_cache(cache, preserve=files[-1])

    assert set(cache.iterdir()) == {files[-1], files[-2]}


def test_recent_document_preview_cache_is_lru_bounded(tmp_path, monkeypatch, qapp):
    import sp.app.folder_navigator.window as window_module

    monkeypatch.setattr(Path, "home", lambda: tmp_path)
    monkeypatch.setattr(window_module, "DOCUMENT_PREVIEW_CACHE_ENTRIES", 2)
    monkeypatch.setattr(
        window_module, "DOCUMENT_PREVIEW_CACHE_BYTES", 1024 * 1024
    )
    paths = [tmp_path / f"recent-{index}.docx" for index in range(3)]
    for path in paths:
        path.write_bytes(path.name.encode())
    window = window_module.Window(tmp_path)
    signatures = [documents.document_signature(path) for path in paths]
    previews = [
        documents.DocumentPreview(backend="python-docx", html=f"<p>{index}</p>")
        for index in range(3)
    ]

    window._cache_document_preview(paths[0], signatures[0], previews[0])
    window._cache_document_preview(paths[1], signatures[1], previews[1])
    assert window._cached_document_preview(paths[0], signatures[0]) is previews[0]
    window._cache_document_preview(paths[2], signatures[2], previews[2])

    assert window._cached_document_preview(paths[1], signatures[1]) is None
    assert window._cached_document_preview(paths[0], signatures[0]) is previews[0]
    assert window._cached_document_preview(paths[2], signatures[2]) is previews[2]
    assert len(window.document_preview_cache) == 2
    window.close()


@pytest.mark.parametrize("suffix", [".docx", ".pptx"])
def test_recent_document_preview_bypasses_rebuild_until_source_changes(
        tmp_path, monkeypatch, qapp, suffix):
    import sp.app.folder_navigator.window as window_module

    path = tmp_path / f"large{suffix}"
    path.write_bytes(b"first version")
    monkeypatch.setattr(Path, "home", lambda: tmp_path)
    builder_name = "build_pptx_preview" if suffix == ".pptx" else "build_docx_preview"
    calls = []

    def build(selected):
        calls.append(documents.document_signature(selected))
        return documents.DocumentPreview(
            backend="python-pptx" if suffix == ".pptx" else "python-docx",
            html=f"<p>render {len(calls)}</p>",
        )

    monkeypatch.setattr(documents, builder_name, build)
    window = window_module.Window(tmp_path)

    def open_and_wait(expected_text):
        window.open_file(path, pinned=True)
        deadline = time.monotonic() + 3
        while True:
            assert time.monotonic() < deadline
            qapp.processEvents()
            tab = window.active_tab()
            viewer = getattr(tab, "viewer", None) if tab is not None else None
            browser = viewer.findChild(window_module.QTextBrowser) if viewer else None
            if browser is not None and browser.toPlainText() == expected_text:
                break
            time.sleep(.01)

    def close_tab():
        tab = window.active_tab()
        window.tabs.removeTab(window.tabs.indexOf(tab))
        tab.deleteLater()
        qapp.processEvents()

    open_and_wait("render 1")
    assert len(calls) == 1
    close_tab()

    open_and_wait("render 1")
    assert len(calls) == 1
    assert window.preview_metrics["document"]["cache_hits"] == 1
    close_tab()

    path.write_bytes(b"second version with a different size")
    open_and_wait("render 2")
    assert len(calls) == 2
    assert window.preview_metrics["document"]["cache_misses"] == 2
    window.close()


@pytest.mark.parametrize("suffix", [".docx", ".pptx"])
def test_libreoffice_preview_can_switch_to_selectable_text(
        tmp_path, monkeypatch, qapp, suffix):
    from PySide6.QtCore import Qt
    from PySide6.QtWidgets import QPushButton, QTextBrowser, QWidget
    import sp.app.folder_navigator.window as window_module

    path = tmp_path / f"selectable{suffix}"
    path.write_bytes(b"placeholder")
    pdf = tmp_path / "preview.pdf"
    pdf.write_bytes(b"%PDF placeholder")
    monkeypatch.setattr(Path, "home", lambda: tmp_path)
    builder_name = "build_pptx_preview" if suffix == ".pptx" else "build_docx_preview"
    monkeypatch.setattr(
        documents,
        builder_name,
        lambda _path: documents.DocumentPreview(
            backend="libreoffice",
            pdf_path=pdf,
            html="<p>Select this document text</p>",
        ),
    )

    layout_views = []

    def fake_pdf_view(_path, **_kwargs):
        widget = QWidget()
        widget.setProperty("folderNavigatorFocusTarget", True)
        widget.zoom_in = lambda: None
        widget.zoom_out = lambda: None
        layout_views.append(widget)
        return widget

    monkeypatch.setattr(window_module, "pdf_view", fake_pdf_view)
    window = window_module.Window(tmp_path)
    window.show()
    window.open_file(path, pinned=True)

    deadline = time.monotonic() + 2
    selectable = None
    while selectable is None:
        assert time.monotonic() < deadline
        qapp.processEvents()
        viewer = getattr(window.active_tab(), "viewer", None)
        selectable = viewer.findChild(QTextBrowser) if viewer is not None else None
        time.sleep(.01)
    buttons = {
        button.text(): button
        for button in window.active_tab().viewer.findChildren(QPushButton)
    }
    buttons["Selectable Text"].click()

    assert selectable.toPlainText() == "Select this document text"
    assert selectable.textInteractionFlags() & Qt.TextSelectableByMouse
    viewer = window.active_tab().viewer
    assert viewer.document_views.currentWidget() is viewer.selectable_text_view
    assert selectable.hasFocus()
    window._focus_tab_content(window.active_tab())
    assert selectable.hasFocus()

    buttons["Layout"].click()
    assert viewer.document_views.currentWidget() is layout_views[0]
    assert layout_views[0].hasFocus()
    window.close()


def test_office_selectable_text_has_pdf_search_focus_and_scroll_parity(qapp):
    from PySide6.QtCore import Qt
    from PySide6.QtTest import QTest
    from sp.app.folder_navigator.window import document_text_view

    paragraphs = [f"Line {number} with needle" for number in range(120)]
    preview = document_text_view(
        "".join(f"<p>{line}</p>" for line in paragraphs), vi_enabled=True
    )
    preview.resize(600, 300)
    preview.show()
    preview.activateWindow()
    qapp.processEvents()
    browser = preview.text_browser
    browser.setFocus()

    QTest.keyClick(browser, Qt.Key_F, Qt.ControlModifier)
    assert preview.document_search.hasFocus()
    assert not preview.document_search_bar.isHidden()
    preview.document_search.setText("needle")
    qapp.processEvents()
    assert preview.document_search_status.text() == "1 of 120 · Start"
    assert not preview.document_previous_match.isEnabled()
    assert preview.document_next_match.isEnabled()

    preview.document_next_match.click()
    assert preview.document_search_status.text() == "2 of 120"
    assert browser.textCursor().selectedText() == "needle"
    preview.document_search.setText("Line 119")
    assert preview.document_search_status.text() == "1 of 1 · Only match"

    QTest.keyClick(preview.document_search, Qt.Key_Escape)
    qapp.processEvents()
    assert preview.document_search_bar.isHidden()
    assert preview.document_search.text() == ""
    assert browser.hasFocus()

    scrollbar = browser.verticalScrollBar()
    scrollbar.setValue(0)
    page_step = scrollbar.pageStep()
    QTest.keyClick(browser, Qt.Key_Down)
    down_value = scrollbar.value()
    assert 0 < down_value < page_step
    QTest.keyClick(browser, Qt.Key_J)
    assert down_value < scrollbar.value() < page_step
    QTest.keyClick(browser, Qt.Key_K)
    assert scrollbar.value() < page_step

    QTest.keyClick(browser, Qt.Key_Slash)
    assert preview.document_search.hasFocus()
    preview.close()


@pytest.mark.parametrize("suffix", [".docx", ".pptx"])
def test_libreoffice_layout_preview_preserves_shift_enter_focus_after_loading(
        tmp_path, monkeypatch, qapp, suffix):
    from PySide6.QtCore import Qt
    from PySide6.QtGui import QPainter, QPdfWriter
    from PySide6.QtPdfWidgets import QPdfView
    from PySide6.QtTest import QTest
    from PySide6.QtWidgets import QWidget
    from sp.app.folder_navigator.window import Window

    path = tmp_path / f"focus{suffix}"
    path.write_bytes(b"placeholder")
    pdf = tmp_path / "converted.pdf"
    writer = QPdfWriter(str(pdf))
    painter = QPainter(writer)
    for page in range(3):
        painter.drawText(100, 100, f"needle page {page + 1}")
        if page < 2:
            writer.newPage()
    painter.end()
    monkeypatch.setattr(Path, "home", lambda: tmp_path)
    builder_name = "build_pptx_preview" if suffix == ".pptx" else "build_docx_preview"
    monkeypatch.setattr(
        documents,
        builder_name,
        lambda _path: documents.DocumentPreview(
            backend="libreoffice",
            pdf_path=pdf,
            html="<p>needle selectable text</p>",
        ),
    )

    window = Window(tmp_path)
    window.resize(900, 500)
    window.tree.vi_enabled = True
    window.show()
    window._suppress_tree_preview = True
    window.tree.setCurrentIndex(window.model.index(str(path)))
    window._suppress_tree_preview = False
    window.tree.setFocus()
    QTest.keyClick(window.tree, Qt.Key_Return, Qt.ShiftModifier)

    deadline = time.monotonic() + 3
    pdf_viewer = None
    while pdf_viewer is None or not pdf_viewer.hasFocus():
        assert time.monotonic() < deadline
        qapp.processEvents()
        tab = window.active_tab()
        viewer = getattr(tab, "viewer", None) if tab is not None else None
        pdf_viewer = viewer.findChild(QPdfView) if viewer is not None else None
        QTest.qWait(10)

    assert pdf_viewer.hasFocus()
    QTest.keyClick(pdf_viewer, Qt.Key_Slash)
    qapp.processEvents()
    pdf_container = next(
        child for child in window.active_tab().viewer.findChildren(QWidget)
        if hasattr(child, "pdf_search")
    )
    assert pdf_container.pdf_search.hasFocus()
    QTest.keyClick(pdf_container.pdf_search, Qt.Key_Escape)
    qapp.processEvents()
    assert pdf_viewer.hasFocus()
    window.close()
