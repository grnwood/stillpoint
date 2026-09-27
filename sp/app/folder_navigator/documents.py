"""Portable, bounded previews for office documents and presentations."""

from __future__ import annotations

import base64
from dataclasses import dataclass
import hashlib
import html
import os
from pathlib import Path
import shutil
import subprocess
import tempfile


MAX_DOCX_BYTES = 100 * 1024 * 1024
MAX_PPTX_BYTES = 200 * 1024 * 1024
MAX_INLINE_IMAGE_BYTES = 12 * 1024 * 1024
MAX_OFFICE_CACHE_BYTES = 128 * 1024 * 1024
MAX_OFFICE_CACHE_FILES = 100
MAX_OFFICE_CACHE_AGE_SECONDS = 30 * 24 * 60 * 60


@dataclass(frozen=True)
class DocumentPreview:
    backend: str
    pdf_path: Path | None = None
    html: str = ""
    notice: str = ""


DocxPreview = DocumentPreview


def document_signature(path: Path) -> tuple[int, int]:
    stat = Path(path).stat()
    return stat.st_mtime_ns, stat.st_size


docx_signature = document_signature


def _cache_key(path: Path) -> str:
    modified, size = document_signature(path)
    material = f"{Path(path).resolve()}\0{modified}\0{size}".encode("utf-8")
    return hashlib.sha256(material).hexdigest()


def _libreoffice_executable() -> str | None:
    return shutil.which("libreoffice") or shutil.which("soffice")


def _trim_office_cache(cache_dir: Path, preserve: Path | None = None) -> None:
    """Bound generated previews by age, count, and total disk usage."""
    import time

    try:
        files = [
            path for path in cache_dir.iterdir()
            if path.is_file() and path.suffix.casefold() in {".pdf", ".tmp"}
        ]
    except OSError:
        return
    cutoff = time.time() - MAX_OFFICE_CACHE_AGE_SECONDS
    retained = []
    for path in files:
        try:
            stat = path.stat()
            if path != preserve and stat.st_mtime < cutoff:
                path.unlink(missing_ok=True)
            else:
                retained.append((path, stat.st_mtime, stat.st_size))
        except OSError:
            continue
    retained.sort(key=lambda item: item[1], reverse=True)
    total = sum(item[2] for item in retained)
    for index, (path, _modified, size) in enumerate(retained):
        if path == preserve:
            continue
        if index >= MAX_OFFICE_CACHE_FILES or total > MAX_OFFICE_CACHE_BYTES:
            try:
                path.unlink(missing_ok=True)
                total -= size
            except OSError:
                pass


def _convert_with_libreoffice(path: Path, executable: str, cache_dir: Path) -> Path:
    cache_dir.mkdir(parents=True, exist_ok=True)
    cached = cache_dir / f"{_cache_key(path)}.pdf"
    _trim_office_cache(cache_dir, preserve=cached)
    if cached.is_file() and cached.stat().st_size:
        try:
            cached.touch()
        except OSError:
            pass
        return cached

    with tempfile.TemporaryDirectory(prefix="stillpoint-office-") as temporary:
        work = Path(temporary)
        output_dir = work / "output"
        profile_dir = work / "profile"
        home_dir = work / "home"
        config_dir = work / "config"
        cache_home = work / "cache"
        runtime_dir = work / "runtime"
        output_dir.mkdir()
        profile_dir.mkdir()
        for directory in (home_dir, config_dir, cache_home, runtime_dir):
            directory.mkdir()
        runtime_dir.chmod(0o700)
        environment = os.environ.copy()
        environment.update({
            "HOME": str(home_dir),
            "XDG_CONFIG_HOME": str(config_dir),
            "XDG_CACHE_HOME": str(cache_home),
            "XDG_RUNTIME_DIR": str(runtime_dir),
        })
        command = [
            executable,
            "--headless",
            "--nologo",
            "--nodefault",
            "--nolockcheck",
            "--norestore",
            f"-env:UserInstallation={profile_dir.as_uri()}",
            "--convert-to",
            "pdf:writer_pdf_Export",
            "--outdir",
            str(output_dir),
            str(path),
        ]
        completed = subprocess.run(
            command,
            capture_output=True,
            timeout=45,
            check=False,
            env=environment,
        )
        converted = output_dir / f"{path.stem}.pdf"
        if completed.returncode or not converted.is_file() or not converted.stat().st_size:
            detail = completed.stderr.decode("utf-8", errors="replace").strip()
            raise RuntimeError(detail or "LibreOffice did not produce a PDF")
        staged = cached.with_suffix(".tmp")
        shutil.copyfile(converted, staged)
        staged.replace(cached)
    _trim_office_cache(cache_dir, preserve=cached)
    return cached


def _runs_html(paragraph) -> str:
    pieces: list[str] = []
    for run in paragraph.runs:
        value = html.escape(run.text).replace("\n", "<br>")
        if not value:
            continue
        if run.bold:
            value = f"<strong>{value}</strong>"
        if run.italic:
            value = f"<em>{value}</em>"
        if run.underline:
            value = f"<u>{value}</u>"
        pieces.append(value)
    return "".join(pieces) or html.escape(paragraph.text)


def _paragraph_html(paragraph) -> str:
    content = _runs_html(paragraph)
    if not content.strip():
        return "<p><br></p>"
    style = (getattr(paragraph.style, "name", "") or "").strip()
    lowered = style.casefold()
    if lowered.startswith("heading"):
        digits = "".join(character for character in style if character.isdigit())
        level = max(1, min(6, int(digits or "1")))
        return f"<h{level}>{content}</h{level}>"
    if "list bullet" in lowered:
        return f"<p class='list'>•&nbsp; {content}</p>"
    if "list number" in lowered:
        return f"<p class='list'>◦&nbsp; {content}</p>"
    return f"<p>{content}</p>"


def _table_html(table) -> str:
    rows = []
    for row_number, row in enumerate(table.rows):
        cells = []
        tag = "th" if row_number == 0 else "td"
        for cell in row.cells:
            content = "<br>".join(
                _runs_html(paragraph) for paragraph in cell.paragraphs
            )
            cells.append(f"<{tag}>{content}</{tag}>")
        rows.append(f"<tr>{''.join(cells)}</tr>")
    return f"<table>{''.join(rows)}</table>"


def _inline_images_html(document) -> str:
    images = []
    total = 0
    seen = set()
    for relationship in document.part.rels.values():
        if not str(relationship.reltype).endswith("/image"):
            continue
        part = relationship.target_part
        blob = bytes(part.blob)
        digest = hashlib.sha256(blob).digest()
        if digest in seen or total + len(blob) > MAX_INLINE_IMAGE_BYTES:
            continue
        seen.add(digest)
        total += len(blob)
        mime = getattr(part, "content_type", "image/png") or "image/png"
        encoded = base64.b64encode(blob).decode("ascii")
        images.append(f"<img src='data:{html.escape(mime)};base64,{encoded}'>")
    if not images:
        return ""
    return "<section class='images'><h2>Images</h2>" + "".join(images) + "</section>"


def _python_docx_html(path: Path) -> str:
    from docx import Document
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(str(path))
    body = []
    for child in document.element.body.iterchildren():
        if isinstance(child, CT_P):
            body.append(_paragraph_html(Paragraph(child, document._body)))
        elif isinstance(child, CT_Tbl):
            body.append(_table_html(Table(child, document._body)))
    body.append(_inline_images_html(document))
    return """<!doctype html><html><head><meta charset='utf-8'><style>
body { font-family: sans-serif; line-height: 1.45; margin: 28px; }
h1,h2,h3,h4,h5,h6 { margin: 1.15em 0 .45em; }
p { margin: .55em 0; } .list { margin-left: 1.5em; }
table { border-collapse: collapse; margin: 1em 0; width: 100%; }
th,td { border: 1px solid #999; padding: 6px 8px; text-align: left; vertical-align: top; }
th { background: rgba(127,127,127,.18); }
img { display: block; max-width: 100%; height: auto; margin: 12px 0; }
</style></head><body>""" + "".join(body) + "</body></html>"


def build_docx_preview(path: Path, cache_dir: Path | None = None) -> DocumentPreview:
    path = Path(path)
    if path.stat().st_size > MAX_DOCX_BYTES:
        raise ValueError(f"DOCX exceeds the {MAX_DOCX_BYTES // (1024 * 1024)} MB preview limit")
    executable = _libreoffice_executable()
    if executable:
        try:
            destination = cache_dir or (Path.home() / ".stillpoint_cache" / "docx")
            pdf_path = _convert_with_libreoffice(path, executable, destination)
            try:
                selectable_html = _python_docx_html(path)
            except Exception:
                selectable_html = ""
            return DocumentPreview(
                backend="libreoffice",
                pdf_path=pdf_path,
                # Keep an accessible/selectable representation alongside the
                # high-fidelity PDF canvas. The UI exposes it on demand.
                html=selectable_html,
            )
        except (OSError, subprocess.SubprocessError, RuntimeError) as exc:
            notice = f"LibreOffice conversion failed; showing simplified preview: {exc}"
            return DocumentPreview(
                backend="python-docx",
                html=_python_docx_html(path),
                notice=notice,
            )
    return DocumentPreview(
        backend="python-docx",
        html=_python_docx_html(path),
        notice="Simplified preview — install LibreOffice for print-layout fidelity.",
    )


def _pptx_paragraph_html(paragraph) -> str:
    pieces = []
    for run in paragraph.runs:
        value = html.escape(run.text).replace("\n", "<br>")
        if not value:
            continue
        font = run.font
        if font.bold:
            value = f"<strong>{value}</strong>"
        if font.italic:
            value = f"<em>{value}</em>"
        if font.underline:
            value = f"<u>{value}</u>"
        pieces.append(value)
    content = "".join(pieces) or html.escape(paragraph.text)
    if not content.strip():
        return ""
    indent = max(0, int(getattr(paragraph, "level", 0) or 0)) * 1.4
    return f"<p style='margin-left:{indent:.1f}em'>{content}</p>"


def _pptx_text_frame_html(text_frame) -> str:
    return "".join(
        _pptx_paragraph_html(paragraph) for paragraph in text_frame.paragraphs
    )


def _pptx_table_html(table) -> str:
    rows = []
    for row_number, row in enumerate(table.rows):
        tag = "th" if row_number == 0 else "td"
        cells = []
        for cell in row.cells:
            cells.append(f"<{tag}>{_pptx_text_frame_html(cell.text_frame)}</{tag}>")
        rows.append(f"<tr>{''.join(cells)}</tr>")
    return f"<table>{''.join(rows)}</table>"


def _pptx_picture_html(shape, image_state: dict) -> str:
    try:
        blob = bytes(shape.image.blob)
        mime = shape.image.content_type or "image/png"
    except Exception:
        return ""
    digest = hashlib.sha256(blob).digest()
    if digest in image_state["seen"]:
        return ""
    if image_state["bytes"] + len(blob) > MAX_INLINE_IMAGE_BYTES:
        return "<p class='unsupported'>Image omitted from simplified preview.</p>"
    image_state["seen"].add(digest)
    image_state["bytes"] += len(blob)
    encoded = base64.b64encode(blob).decode("ascii")
    return f"<img src='data:{html.escape(mime)};base64,{encoded}'>"


def _pptx_shape_html(shape, image_state: dict) -> str:
    if getattr(shape, "has_table", False):
        return _pptx_table_html(shape.table)
    if getattr(shape, "shape_type", None) == 13:  # MSO_SHAPE_TYPE.PICTURE
        return _pptx_picture_html(shape, image_state)
    if getattr(shape, "has_chart", False):
        chart = shape.chart
        title = "Chart"
        try:
            if chart.has_title and chart.chart_title.has_text_frame:
                title = chart.chart_title.text_frame.text or title
        except Exception:
            pass
        series = [html.escape(str(item.name)) for item in chart.series if item.name]
        detail = f": {', '.join(series)}" if series else ""
        return f"<p class='unsupported'><strong>{html.escape(title)}</strong>{detail}</p>"
    if getattr(shape, "shape_type", None) == 6:  # MSO_SHAPE_TYPE.GROUP
        children = sorted(shape.shapes, key=lambda item: (item.top, item.left))
        return "".join(_pptx_shape_html(child, image_state) for child in children)
    if getattr(shape, "has_text_frame", False):
        return _pptx_text_frame_html(shape.text_frame)
    return ""


def _python_pptx_html(path: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(str(path))
    image_state = {"seen": set(), "bytes": 0}
    slides = []
    for number, slide in enumerate(presentation.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None else ""
        title = title or f"Slide {number}"
        blocks = []
        shapes = sorted(slide.shapes, key=lambda item: (item.top, item.left))
        for shape in shapes:
            if title_shape is not None and shape._element is title_shape._element:
                continue
            rendered = _pptx_shape_html(shape, image_state)
            if rendered:
                blocks.append(rendered)
        notes = ""
        try:
            notes_text = slide.notes_slide.notes_text_frame.text.strip()
            if notes_text:
                notes = (
                    "<div class='notes'><strong>Speaker notes</strong>"
                    f"<p>{html.escape(notes_text).replace(chr(10), '<br>')}</p></div>"
                )
        except Exception:
            pass
        slides.append(
            "<section class='slide'>"
            f"<div class='slide-number'>Slide {number}</div>"
            f"<h1>{html.escape(title)}</h1>"
            f"{''.join(blocks) or '<p class=\"unsupported\">No extractable slide content.</p>'}"
            f"{notes}</section>"
        )
    return """<!doctype html><html><head><meta charset='utf-8'><style>
body { font-family: sans-serif; margin: 24px; background: #d9d9d9; color: #111; }
.slide { background: #fff; border: 1px solid #999; margin: 0 auto 24px; padding: 28px;
         max-width: 960px; min-height: 420px; }
.slide-number { color: #666; font-size: 12px; text-align: right; }
h1 { font-size: 28px; margin: 0 0 20px; } p { margin: .55em 0; }
table { border-collapse: collapse; margin: 14px 0; width: 100%; }
th,td { border: 1px solid #999; padding: 6px 8px; text-align: left; vertical-align: top; }
th { background: #eee; } img { display: block; max-width: 100%; height: auto; margin: 12px auto; }
.unsupported { color: #555; font-style: italic; }
.notes { border-top: 1px solid #bbb; color: #444; margin-top: 24px; padding-top: 12px; }
</style></head><body>""" + "".join(slides) + "</body></html>"


def build_pptx_preview(path: Path, cache_dir: Path | None = None) -> DocumentPreview:
    path = Path(path)
    if path.stat().st_size > MAX_PPTX_BYTES:
        raise ValueError(f"PPTX exceeds the {MAX_PPTX_BYTES // (1024 * 1024)} MB preview limit")
    executable = _libreoffice_executable()
    if executable:
        try:
            destination = cache_dir or (Path.home() / ".stillpoint_cache" / "pptx")
            pdf_path = _convert_with_libreoffice(path, executable, destination)
            try:
                selectable_html = _python_pptx_html(path)
            except Exception:
                selectable_html = ""
            return DocumentPreview(
                backend="libreoffice",
                pdf_path=pdf_path,
                html=selectable_html,
            )
        except (OSError, subprocess.SubprocessError, RuntimeError) as exc:
            return DocumentPreview(
                backend="python-pptx",
                html=_python_pptx_html(path),
                notice=f"LibreOffice conversion failed; showing simplified preview: {exc}",
            )
    return DocumentPreview(
        backend="python-pptx",
        html=_python_pptx_html(path),
        notice="Simplified preview — install LibreOffice for slide-layout fidelity.",
    )
